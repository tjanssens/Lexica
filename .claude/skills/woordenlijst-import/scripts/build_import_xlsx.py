#!/usr/bin/env python3
"""Bouw een Lexica-importeerbaar Excel-bestand uit een JSON-woordenlijst.

Gebruik: python3 build_import_xlsx.py <invoer.json> <uitvoer.xlsx>

Het JSON-formaat staat beschreven in ../SKILL.md.
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "number", "language*", "term*", "translation*", "part_of_speech",
    "notes", "easiness", "interval", "repetitions", "due_date", "group",
]
WIDTHS = [8, 12, 32, 32, 22, 55, 10, 10, 12, 12, 26]

# Zelfde aliassen als ExcelImportService.LanguageAliases.
LANGUAGES = {
    "french": "French", "fr": "French", "frans": "French",
    "english": "English", "en": "English", "engels": "English",
    "latin": "Latin", "la": "Latin", "latijn": "Latin",
    "greek": "Greek", "el": "Greek", "grieks": "Greek",
}


def build(data, out_path):
    default_language = data.get("language")
    default_group = data.get("group", "")
    words = data.get("words") or []
    if not words:
        sys.exit("Geen woorden gevonden in de invoer (verwacht een lijst onder 'words').")

    wb = Workbook()
    ws = wb.active
    ws.title = "Words"

    fill = PatternFill("solid", fgColor="0F3460")
    font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = fill
        cell.font = font

    seen = {}
    for index, word in enumerate(words, start=1):
        term = (word.get("term") or "").strip()
        translation = (word.get("translation") or "").strip()
        if not term:
            sys.exit(f"Rij {index}: 'term' ontbreekt.")
        if not translation:
            sys.exit(f"Rij {index}: 'translation' ontbreekt voor '{term}'.")

        raw_language = word.get("language") or default_language
        if not raw_language:
            sys.exit(f"Rij {index}: geen taal opgegeven en geen standaardtaal in de invoer.")
        language = LANGUAGES.get(str(raw_language).strip().lower())
        if language is None:
            sys.exit(f"Rij {index}: onbekende taal '{raw_language}'. "
                     f"Kies uit: Latin, Greek, English, French.")

        key = (language, term.lower())
        if key in seen:
            print(f"Let op: '{term}' ({language}) staat al op rij {seen[key]}; "
                  f"de import markeert rij {index} als duplicaat.", file=sys.stderr)
        else:
            seen[key] = index

        row = index + 1
        ws.cell(row, 1, index)
        ws.cell(row, 2, language)
        ws.cell(row, 3, term)
        ws.cell(row, 4, translation)
        ws.cell(row, 5, word.get("partOfSpeech", "") or "")
        ws.cell(row, 6, word.get("notes", "") or "")
        ws.cell(row, 7, word.get("easiness", 2.5))
        ws.cell(row, 8, word.get("interval", 0))
        ws.cell(row, 9, word.get("repetitions", 0))
        # due_date bewust leeg laten tenzij expliciet gevraagd: de import vult dan zelf
        # vandaag in (in UTC). Een datum uit Excel komt zonder tijdzone binnen en werd
        # door oudere serverversies geweigerd door Postgres.
        ws.cell(row, 10, word.get("dueDate", "") or "")
        ws.cell(row, 11, word.get("group", default_group) or "")

    for col, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    wb.save(out_path)
    print(f"{len(words)} woorden weggeschreven naar {out_path}")


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as handle:
        build(json.load(handle), sys.argv[2])


if __name__ == "__main__":
    main()
